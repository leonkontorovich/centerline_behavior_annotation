import pandas as pd
import os
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
import os
import argparse
import sys

def rev_duration(input_path, project_folder, output_file, fps):
    # Loading the dataset
    data = pd.read_csv(input_path, header=None)
    data.columns = ['frames', 'values']

    # Identifying reversals in the data
    reversals = data['values'] == -1

    # Finding starts and ends of reversals in the data
    starts = data['frames'][reversals & (~reversals.shift(1, fill_value=False))]
    ends = data['frames'][reversals & (~reversals.shift(-1, fill_value=False))]

    reversal_stats = [(f"Reversal {i}", start, end, (end - start) / fps)
                      for i, (start, end) in enumerate(zip(starts, ends), start=1)]

    # Creating a DataFrame for the reversal statistics
    reversal_stats_df = pd.DataFrame(reversal_stats, columns=["Reversal", "Start Frame", "End Frame", "Duration (seconds)"])

    # Creating the Excel file with centered alignment and autofit columns
    wb = openpyxl.Workbook()
    ws = wb.active

    for r_idx, row in enumerate(dataframe_to_rows(reversal_stats_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.alignment = openpyxl.styles.Alignment(horizontal='center')

    for column_cells in ws.columns:
        length = max(len(str(cell.value)) for cell in column_cells)
        ws.column_dimensions[openpyxl.utils.get_column_letter(column_cells[0].column)].width = length


    # Save the workbook
    wb.save(output_file)
    print(f"Modified Excel file saved at: {output_file}")


def main(arg_list=None):
    parser = argparse.ArgumentParser(description='Rev Duration')
    parser.add_argument('--input_path', help='folder with the tracker position', required=True)
    parser.add_argument('--beh_annotation1', help='Path to CSV with Behavior Annotations', required=True)
    parser.add_argument('--output_file', help='Output File Name', required=True)
    parser.add_argument('--fps', help='FPS', required=True)

    args = parser.parse_args(arg_list)

    main_folder = args.input_path
    etho_path1 = args.beh_annotation1
    output_file = args.output_file
    fps = int(args.fps)

    project_folder = main_folder
    print("Project folder is: ", project_folder)

    rev_duration(etho_path1, project_folder, output_file, fps)


if __name__ == '__main__':
    main(sys.argv[1:])

